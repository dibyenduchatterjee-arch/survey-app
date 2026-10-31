from pathlib import Path

import pytest
from sqlalchemy import func, select

from app.core.security import create_access_token
from app.db.models import Assignment, AuditLog, Dataset, Form, FormVersion

FIX = Path(__file__).parent / "fixtures"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _files(name="amb_asha_visit.xlsx", data=None):
    data = (FIX / name).read_bytes() if data is None else data
    return {"file": (name, data, XLSX)}


@pytest.fixture
def as_(user_factory):
    def make(role, username=None):
        u = user_factory(username=username or role, role=role)
        return u, {"Authorization": f"Bearer {create_access_token(u.id, role)}"}

    return make


@pytest.fixture
def sup(as_):
    return as_("supervisor")


def _upload(client, headers, name="amb_asha_visit.xlsx"):
    r = client.post("/api/forms", files=_files(name), headers=headers)
    assert r.status_code == 201, r.text
    return r.json()


def _publish(client, headers, form_id, version):
    return client.post(
        f"/api/forms/{form_id}/publish", json={"version": version}, headers=headers
    )


def _count(db, model):
    return db.scalar(select(func.count()).select_from(model))


def test_supervisor_can_upload_and_publish(client, db, sup):
    user, h = sup
    body = _upload(client, h)
    assert set(body) == {"form_id", "version", "warnings"}
    assert body["warnings"] == []
    assert len(body["version"]) == 10 and body["version"].isdigit()
    assert body["version"].endswith("01")
    r = _publish(client, h, body["form_id"], body["version"])
    assert r.status_code == 200
    assert r.json() == {
        "form_id": body["form_id"], "version": body["version"], "is_current": True
    }
    v = db.scalars(select(FormVersion)).one()
    assert v.is_current and v.published_by == user.id


def test_unauthenticated_upload_is_401(client):
    assert client.post("/api/forms", files=_files()).status_code == 401


def test_enumerator_cannot_upload(client, as_):
    _, h = as_("enumerator")
    assert client.post("/api/forms", files=_files(), headers=h).status_code == 403


def test_enumerator_cannot_publish_or_upload_dataset(client, as_, sup):
    _, sh = sup
    body = _upload(client, sh)
    _, eh = as_("enumerator")
    assert _publish(client, eh, body["form_id"], body["version"]).status_code == 403
    r = client.post("/api/datasets/blocks", files={"file": ("b.csv", b"name\nx\n")}, headers=eh)
    assert r.status_code == 403


def _xlsx_with_bad_age_constraint():
    import io

    import openpyxl

    wb = openpyxl.load_workbook(FIX / "amb_asha_visit.xlsx")
    ws = wb["survey"]
    header = [c.value for c in ws[1]]
    col = header.index("constraint") + 1
    for row in ws.iter_rows(min_row=2):
        if row[header.index("name")].value == "child_age_months":
            ws.cell(row=row[0].row, column=col).value = "substr(., 0, 1) = '1'"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _assert_rejected_cleanly(db):
    assert _count(db, FormVersion) == 0
    assert _count(db, Form) == 0
    assert _count(db, AuditLog) == 0


def test_unsupported_expression_returns_422_names_the_question_and_leaves_nothing(
    client, db, sup
):
    _, h = sup
    r = client.post("/api/forms", files=_files("amb_unsupported_expr.xlsx"), headers=h)
    assert r.status_code == 422
    detail = r.json()["detail"]
    assert detail["error"] == "unsupported_expression"
    assert detail["question"] == "initial" and "initial" in r.text  # that fixture's offender
    _assert_rejected_cleanly(db)


def test_the_422_names_child_age_months_when_that_is_the_offender(client, db, sup):
    _, h = sup
    r = client.post(
        "/api/forms", files=_files("age.xlsx", _xlsx_with_bad_age_constraint()), headers=h
    )
    assert r.status_code == 422 and "child_age_months" in r.text
    assert r.json()["detail"]["question"] == "child_age_months"
    _assert_rejected_cleanly(db)


def test_a_failed_second_upload_leaves_the_first_intact(client, db, sup):
    _, h = sup
    _upload(client, h)
    r = client.post("/api/forms", files=_files("broken_missing_choices.xlsx"), headers=h)
    assert r.status_code == 422
    assert _count(db, FormVersion) == 1


def test_conversion_failure_is_422(client, db, sup):
    _, h = sup
    r = client.post("/api/forms", files=_files("broken_missing_choices.xlsx"), headers=h)
    assert r.status_code == 422
    assert _count(db, FormVersion) == 0


def test_non_xlsx_bytes_are_refused_whatever_the_client_claims(client, db, sup):
    _, h = sup
    r = client.post("/api/forms", files=_files("evil.xlsx", b"%PDF-1.7 not a zip"), headers=h)
    assert r.status_code == 415
    assert _count(db, FormVersion) == 0


def test_oversized_upload_is_413(client, db, sup):
    _, h = sup
    big = b"PK\x03\x04" + b"0" * (5 * 1024 * 1024)
    r = client.post("/api/forms", files=_files("big.xlsx", big), headers=h)
    assert r.status_code == 413
    assert _count(db, FormVersion) == 0


def test_versions_are_date_plus_sequence(client, sup):
    _, h = sup
    v1 = _upload(client, h)["version"]
    v2 = _upload(client, h)["version"]
    assert v1[:8] == v2[:8]
    assert (v1[8:], v2[8:]) == ("01", "02")


def test_a_second_publish_moves_is_current(client, db, sup):
    _, h = sup
    a = _upload(client, h)
    b = _upload(client, h)
    assert a["form_id"] == b["form_id"] and a["version"] != b["version"]
    assert _publish(client, h, a["form_id"], a["version"]).status_code == 200
    assert _publish(client, h, b["form_id"], b["version"]).status_code == 200
    current = db.scalars(select(FormVersion).where(FormVersion.is_current)).all()
    assert [v.version for v in current] == [b["version"]]
    assert _count(db, FormVersion) == 2


def test_publish_unknown_form_or_version_is_404(client, sup):
    _, h = sup
    body = _upload(client, h)
    assert _publish(client, h, "nope", body["version"]).status_code == 404
    assert _publish(client, h, body["form_id"], "2000010101").status_code == 404


def test_an_old_version_spec_is_still_served_after_a_new_publish(client, as_, db, sup):
    _, h = sup
    a = _upload(client, h)
    _publish(client, h, a["form_id"], a["version"])
    b = _upload(client, h)
    _publish(client, h, b["form_id"], b["version"])
    r = client.get(f"/api/forms/{a['form_id']}/versions/{a['version']}/spec", headers=h)
    assert r.status_code == 200
    spec = r.json()
    assert spec["nodes"] and spec["version"] == a["version"]
    # an assigned enumerator can finish a draft started on the old version
    eu, eh = as_("enumerator")
    db.add(Assignment(user_id=eu.id, form_id=db.scalar(select(Form)).id))
    db.commit()
    r = client.get(f"/api/forms/{a['form_id']}/versions/{a['version']}/spec", headers=eh)
    assert r.status_code == 200


def test_enumerator_sees_only_assigned_forms(client, as_, db, sup):
    su, sh = sup
    body = _upload(client, sh)
    _publish(client, sh, body["form_id"], body["version"])
    db.add(Form(form_id="other_form", title="Other", created_by=su.id))
    db.commit()
    eu, eh = as_("enumerator")
    assert client.get("/api/forms", headers=eh).json() == []
    mine = db.scalar(select(Form).where(Form.form_id == body["form_id"]))
    db.add(Assignment(user_id=eu.id, form_id=mine.id))
    db.commit()
    r = client.get("/api/forms", headers=eh)
    rows = r.json()
    assert [x["form_id"] for x in rows] == [body["form_id"]]
    assert rows[0]["current_version"] == body["version"]
    assert "other_form" not in r.text


def test_enumerator_gets_404_for_an_unassigned_forms_spec(client, as_, sup):
    _, sh = sup
    body = _upload(client, sh)
    _publish(client, sh, body["form_id"], body["version"])
    _, eh = as_("enumerator")
    r = client.get(f"/api/forms/{body['form_id']}/versions/{body['version']}/spec", headers=eh)
    assert r.status_code == 404


def test_supervisor_lists_all_forms(client, sup):
    _, h = sup
    body = _upload(client, h)
    rows = client.get("/api/forms", headers=h).json()
    assert rows[0]["form_id"] == body["form_id"] and rows[0]["current_version"] is None


def test_publish_writes_an_audit_row(client, db, sup):
    user, h = sup
    body = _upload(client, h)
    _publish(client, h, body["form_id"], body["version"])
    row = db.scalars(select(AuditLog).where(AuditLog.action == "form.publish")).one()
    assert row.actor_user_id == user.id
    assert row.object_type == "form" and row.object_id == body["form_id"]
    assert row.detail_json == {"version": body["version"]}


def test_dataset_csv_round_trips_devanagari(client, sup):
    _, h = sup
    r = client.post(
        "/api/datasets/blocks",
        files={"file": ("b.csv", "name\nमुजफ्फरपुर\n".encode("utf-8"))},
        headers=h,
    )
    assert r.status_code == 200
    got = client.get("/api/datasets/blocks", headers=h).json()
    assert got["rows"][0]["name"] == "मुजफ्फरपुर"
    assert got["name"] == "blocks" and got["version"] == 1


def test_dataset_reupload_bumps_version_and_enumerator_can_read(client, as_, sup):
    _, h = sup
    for _i in range(2):
        client.post("/api/datasets/blocks", files={"file": ("b.csv", b"name\nA\n")}, headers=h)
    _, eh = as_("enumerator")
    got = client.get("/api/datasets/blocks", headers=eh)
    assert got.status_code == 200 and got.json()["version"] == 2


def test_dataset_missing_is_404_and_needs_auth(client, sup):
    _, h = sup
    assert client.get("/api/datasets/nope", headers=h).status_code == 404
    assert client.get("/api/datasets/nope").status_code == 401


@pytest.mark.parametrize(
    "name,data",
    [
        ("blocks", b""),
        ("blocks", b"\xff\xfe\x00bad"),
        ("blocks", b"PK\x03\x04zip"),
        ("blocks", b"a,b\n1\n"),
        ("bad name!", b"name\nx\n"),
    ],
)
def test_bad_datasets_are_refused_and_store_nothing(client, db, sup, name, data):
    _, h = sup
    r = client.post(f"/api/datasets/{name}", files={"file": ("b.csv", data)}, headers=h)
    assert r.status_code in (404, 422)
    assert _count(db, Dataset) == 0
