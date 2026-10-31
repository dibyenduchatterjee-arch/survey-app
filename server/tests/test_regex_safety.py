"""regex() patterns that can run for seconds are refused at publish (review finding I-2).

`(a+)+$` takes about 5 s on 26 characters in Python's `re`, and the same pattern runs in the
browser, so a form carrying one would freeze the server and every phone that opens it.
"""

import io

import openpyxl
import pytest

from app.forms.expr import UnsupportedExpression, parse
from test_forms_api import FIX, XLSX, _assert_rejected_cleanly, as_, sup  # noqa: F401 - fixtures

CATASTROPHIC = [
    r"(a+)+$",
    r"^(\d+)*$",
    r"(a*)*b",
    r"([a-z]+ ?)+",
    r"(a|aa)+$",
    r"(\d|\d\d)*",
    r"(a?){30}a{30}",
    r"(?:x+x+)+y",
    r"((ab)*)*c",
    r"(.|a)+",
    r"(?:a+|b)+c",
    r"(.*,)+",
]

FINE = [
    r"[6-9][0-9]{9}",
    r"\d+",
    r"pink|blue",
    r"(\d+\.){3}\d+",        # a dotted number: the outer repeat is small and fixed
    r"(cat|dog)+",           # branches start differently, so nothing is ambiguous
    r"(\d{3}-)+\d{4}",       # a fixed inner count is not a nested quantifier
    r"([A-Z][a-z]{2,8} )+",
    r"[a-z]+(,[a-z]+)*",
    r"(?:pink|blue)(?: syrup)?",
    r"(\d+,){1,3}",
    r"(a|[a-z])+",         # the parser folds this into one character set
]


@pytest.mark.parametrize("pattern", CATASTROPHIC)
def test_a_catastrophic_pattern_is_refused_at_parse_time(pattern):
    with pytest.raises(UnsupportedExpression) as exc:
        parse(f"regex(., '{pattern}')")
    assert "regex()" in str(exc.value) and "backtrack" in str(exc.value)


@pytest.mark.parametrize("pattern", FINE)
def test_ordinary_patterns_still_parse(pattern):
    parse(f"regex(., '{pattern}')")


def _xlsx_with_name_constraint(expr):
    wb = openpyxl.load_workbook(FIX / "amb_asha_visit.xlsx")
    ws = wb["survey"]
    header = [c.value for c in ws[1]]
    col = header.index("constraint") + 1
    for row in ws.iter_rows(min_row=2):
        if row[header.index("name")].value == "child_name":
            ws.cell(row=row[0].row, column=col).value = expr
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_publishing_a_form_with_a_catastrophic_pattern_is_refused_loudly(client, db, sup):
    _, h = sup
    data = _xlsx_with_name_constraint("regex(., '(a+)+$')")
    r = client.post("/api/forms", files={"file": ("f.xlsx", data, XLSX)}, headers=h)
    assert r.status_code == 422
    assert r.json()["detail"]["error"] == "unsupported_expression"
    assert r.json()["detail"]["question"] == "child_name"
    assert "backtrack" in r.text
    _assert_rejected_cleanly(db)
