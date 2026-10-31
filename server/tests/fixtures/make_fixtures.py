"""Builds the XLSForm fixtures used by the form tests (Tasks 7-13, 22).

Run:  .venv/Scripts/python.exe server/tests/fixtures/make_fixtures.py
The .xlsx files are committed so tests need not run this; regenerate after editing.

Fixtures
    amb_asha_visit.xlsx          realistic Anemia Mukt Bihar ASHA visit form. Every
                                 expression stays inside the subset of
                                 server/app/forms/expr.py (see shared/EXPR_CASES.md).
    broken_missing_choices.xlsx  select_one on a list absent from the choices sheet.
    amb_unsupported_expr.xlsx    converts fine but uses an expression the evaluator
                                 rejects (substr); for Task 8's negative test.
"""

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

SURVEY_HEADER = [
    "type", "name", "label::hi (Hindi)", "label::en (English)",
    "hint::hi (Hindi)", "hint::en (English)",
    "required", "relevant", "constraint",
    "constraint_message::hi (Hindi)", "constraint_message::en (English)",
    "calculation", "choice_filter", "appearance",
]

CHOICES_HEADER = ["list_name", "name", "label::hi (Hindi)", "label::en (English)", "district"]

_ALIASES = {
    "label_hi": "label::hi (Hindi)", "label_en": "label::en (English)",
    "hint_hi": "hint::hi (Hindi)", "hint_en": "hint::en (English)",
    "cmsg_hi": "constraint_message::hi (Hindi)",
    "cmsg_en": "constraint_message::en (English)",
}


def _s(**kw):
    out = {_ALIASES.get(k, k): v for k, v in kw.items()}
    unknown = set(out) - set(SURVEY_HEADER)
    assert not unknown, unknown
    return [out.get(h) for h in SURVEY_HEADER]


def _write(path, survey, choices, settings):
    wb = Workbook()
    ws = wb.active
    ws.title = "survey"
    ws.append(SURVEY_HEADER)
    for r in survey:
        ws.append(r)
    ws = wb.create_sheet("choices")
    ws.append(CHOICES_HEADER)
    for r in choices:
        ws.append(r)
    ws = wb.create_sheet("settings")
    ws.append(["form_title", "form_id", "version", "default_language"])
    ws.append(settings)
    wb.save(path)


YES_NO = [
    ["yes_no", "yes", "हाँ", "Yes", None],
    ["yes_no", "no", "नहीं", "No", None],
]

CHOICES = [
    ["district", "patna", "पटना", "Patna", None],
    ["district", "gaya", "गया", "Gaya", None],
    ["district", "muzaffarpur", "मुज़फ़्फ़रपुर", "Muzaffarpur", None],
    ["block", "danapur", "दानापुर", "Danapur", "patna"],
    ["block", "paliganj", "पालीगंज", "Paliganj", "patna"],
    ["block", "bodh_gaya", "बोधगया", "Bodh Gaya", "gaya"],
    ["block", "sherghati", "शेरघाटी", "Sherghati", "gaya"],
    ["block", "motipur", "मोतीपुर", "Motipur", "muzaffarpur"],
    ["block", "kanti", "कांटी", "Kanti", "muzaffarpur"],
    *YES_NO,
    ["sex", "girl", "लड़की", "Girl", None],
    ["sex", "boy", "लड़का", "Boy", None],
    ["ifa_form", "syrup", "आई एफ ए सिरप", "IFA syrup", None],
    ["ifa_form", "pink", "आई एफ ए गुलाबी गोली", "IFA pink tablet", None],
    ["ifa_form", "blue", "आई एफ ए नीली गोली", "IFA blue tablet", None],
]


def amb_asha_visit():
    survey = [
        _s(type="start", name="start"),
        _s(type="end", name="end"),
        _s(type="select_one district", name="district",
           label_hi="ज़िला", label_en="District", required="yes"),
        _s(type="select_one block", name="block", label_hi="प्रखंड",
           label_en="Block", required="yes", choice_filter="district=${district}"),
        _s(type="text", name="asha_name", label_hi="आशा का नाम",
           label_en="ASHA name", required="yes"),
        _s(type="geopoint", name="visit_location",
           label_hi="घर की लोकेशन", label_en="Household location"),
        _s(type="image", name="household_photo", label_hi="घर की फोटो",
           label_en="Household photo", appearance="new"),
        _s(type="integer", name="household_children", required="yes",
           label_hi="घर में पाँच साल से छोटे कितने बच्चे हैं?",
           label_en="How many children under five live in the household?",
           constraint=". >= 0 and . <= 10",
           cmsg_hi="कृपया शून्य से दस के बीच की संख्या डालें",
           cmsg_en="Enter a number from 0 to 10"),
        _s(type="begin repeat", name="children", label_hi="बच्चों की सूची",
           label_en="Children roster"),
        _s(type="text", name="child_name", label_hi="बच्चे का नाम",
           label_en="Child's name", required="yes"),
        _s(type="select_one sex", name="child_sex", label_hi="बच्चे का लिंग",
           label_en="Child's sex"),
        _s(type="integer", name="child_age_months", required="yes",
           label_hi="बच्चे की उम्र (महीनों में)",
           label_en="Child's age (in months)",
           hint_hi="छह से उनसठ महीने", hint_en="6 to 59 months",
           constraint=". >= 6 and . <= 59",
           cmsg_hi="उम्र छह से उनसठ महीने के बीच होनी चाहिए",
           cmsg_en="Age must be between 6 and 59 months"),
        _s(type="calculate", name="child_age_years",
           calculation="int(${child_age_months} div 12)"),
        _s(type="select_one ifa_form", name="ifa_form",
           label_hi="बच्चे को कौन सी आई एफ ए दी जाती है?",
           label_en="Which IFA does the child receive?",
           relevant="${child_age_months} >= 6"),
        _s(type="select_one yes_no", name="syrup_given", required="yes",
           label_hi="क्या बच्चे को आई एफ ए सिरप पिलाया गया?",
           label_en="Was the child given IFA syrup?",
           relevant="${ifa_form} = 'syrup'"),
        _s(type="begin repeat", name="ifa_doses",
           label_hi="आई एफ ए खुराक का रिकॉर्ड", label_en="IFA dose record",
           relevant="${syrup_given} = 'yes'"),
        _s(type="date", name="dose_date", label_hi="खुराक की तारीख",
           label_en="Dose date", required="yes",
           constraint=". <= today()",
           cmsg_hi="तारीख आज से आगे की नहीं हो सकती",
           cmsg_en="The date cannot be in the future"),
        _s(type="integer", name="dose_ml", label_hi="खुराक (मिलीलीटर)",
           label_en="Dose (ml)"),
        _s(type="end repeat"),
        _s(type="end repeat"),
        _s(type="calculate", name="children_count", calculation="count(${child_name})"),
        _s(type="note", name="summary_note",
           label_hi="धन्यवाद, यह दौरा पूरा हुआ", label_en="Thank you, visit complete"),
    ]
    _write(HERE / "amb_asha_visit.xlsx", survey, CHOICES,
           ["Anemia Mukt Bihar - ASHA home visit", "amb_asha_visit",
            "2026103101", "hi (Hindi)"])


def broken_missing_choices():
    survey = [
        _s(type="select_one district", name="district", label_hi="ज़िला",
           label_en="District"),
        _s(type="text", name="note_text", label_hi="टिप्पणी", label_en="Note"),
    ]
    # the choices sheet exists but has no `district` list
    _write(HERE / "broken_missing_choices.xlsx", survey, list(YES_NO),
           ["Broken form", "broken_form", "1", "hi (Hindi)"])


def amb_unsupported_expr():
    survey = [
        _s(type="text", name="asha_name", label_hi="आशा का नाम", label_en="ASHA name"),
        _s(type="calculate", name="initial", calculation="substr(${asha_name}, 0, 1)"),
    ]
    _write(HERE / "amb_unsupported_expr.xlsx", survey, list(YES_NO),
           ["Unsupported expression", "amb_unsupported_expr", "1", "hi (Hindi)"])


if __name__ == "__main__":
    amb_asha_visit()
    broken_missing_choices()
    amb_unsupported_expr()
    print("fixtures written to", HERE)
