import json
from pathlib import Path

import pytest

from app.forms.compile import FormCompileError, compile_spec
from app.forms.expr import UnsupportedExpression
from app.forms.xlsform import to_xform

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module")
def xml():
    return to_xform((FIX / "amb_asha_visit.xlsx").read_bytes())


@pytest.fixture(scope="module")
def spec(xml):
    return compile_spec(xml)


def walk(nodes):
    for n in nodes:
        yield n
        yield from walk(n.get("children", []))


def find(nodes, name):
    return next(n for n in walk(nodes) if n["name"] == name)


def test_spec_carries_form_id_version_and_languages(spec):
    assert spec["form_id"] == "amb_asha_visit"
    assert spec["version"] == "2026103101"
    assert spec["languages"] == ["hi", "en"]
    assert spec["default_language"] == "hi"
    assert spec["title"]["en"] == "Anemia Mukt Bihar - ASHA home visit"


def test_labels_are_keyed_by_short_language_code(spec):
    assert find(spec["nodes"], "district")["label"] == {"hi": "ज़िला", "en": "District"}


def test_question_types_are_mapped(spec):
    types = {n["type"] for n in walk(spec["nodes"])}
    assert types >= {"select_one", "select_multiple", "integer", "geopoint", "image",
                     "repeat", "calculate", "text", "date", "note"}
    assert find(spec["nodes"], "visit_location")["type"] == "geopoint"
    assert find(spec["nodes"], "household_photo")["type"] == "image"
    assert find(spec["nodes"], "symptoms")["type"] == "select_multiple"
    assert find(spec["nodes"], "summary_note")["type"] == "note"


def test_metadata_is_not_a_question(spec):
    names = {n["name"] for n in walk(spec["nodes"])}
    assert not names & {"start", "end", "meta", "instanceID"}


def test_choice_filter_comes_out_in_dollar_form(spec):
    block = find(spec["nodes"], "block")
    assert block["choices_ref"] == "block"
    assert block["choice_filter"] == "district = ${district}"
    assert "choice_filter" not in find(spec["nodes"], "district")


def test_choices_carry_labels_and_filter_properties(spec):
    kanti = next(c for c in spec["choices"]["block"] if c["value"] == "kanti")
    assert kanti["label"] == {"hi": "कांटी", "en": "Kanti"}
    assert kanti["props"] == {"district": "muzaffarpur"}
    assert spec["choices"]["yes_no"][0] == {"value": "yes", "label": {"hi": "हाँ", "en": "Yes"}}


def test_relevant_and_calculate_are_carried_through_in_dollar_form(spec):
    assert find(spec["nodes"], "referral_note")["relevant"] == "selected(${symptoms}, 'pallor')"
    assert find(spec["nodes"], "tablet_taken")["relevant"] == \
        "${ifa_form} != 'syrup' and ${child_age_years} >= 3"
    assert find(spec["nodes"], "child_age_years")["calculate"] == "int(${child_age_months} div 12)"
    assert find(spec["nodes"], "children_count")["calculate"] == "count(${child_name})"
    assert find(spec["nodes"], "ifa_doses")["relevant"] == "${syrup_given} = 'yes'"


def test_constraint_message_keeps_its_languages(spec):
    n = find(spec["nodes"], "child_age_months")
    assert n["constraint"] == ". >= 6 and . <= 59"
    assert n["constraint_message"] == {
        "hi": "उम्र छह से उनसठ महीने के बीच होनी चाहिए",
        "en": "Age must be between 6 and 59 months",
    }
    assert n["hint"]["en"] == "6 to 59 months"
    assert n["required"] is True


def test_a_repeat_inside_a_repeat_compiles_with_both_names(spec):
    outer = find(spec["nodes"], "children")
    inner = find(outer["children"], "ifa_doses")
    assert outer["type"] == inner["type"] == "repeat"
    assert outer["repeat_count"] == "${household_children}"
    assert outer["label"]["hi"] == "बच्चों की सूची"
    assert [c["name"] for c in inner["children"]] == ["dose_date", "dose_ml"]
    # the pyxform wrapper group and the blank template row must not double the nodes
    assert [n["name"] for n in spec["nodes"]].count("children") == 1
    assert [c["name"] for c in outer["children"]].count("ifa_doses") == 1


def test_spec_is_json_with_intact_hindi(spec):
    text = json.dumps(spec, ensure_ascii=False)
    assert "मुज़फ़्फ़रपुर" in text and "???" not in text
    assert json.loads(text) == spec


def test_an_unsupported_expression_fails_compilation_loudly(xml):
    bad = xml.replace(
        ". &gt;= 6 and . &lt;= 59", "instance('ages')/root/item[name=.]/max &gt; 0"
    )
    assert bad != xml
    with pytest.raises(UnsupportedExpression) as e:
        compile_spec(bad)
    assert "child_age_months" in str(e.value)
    assert e.value.question == "child_age_months"


def test_the_unsupported_fixture_names_its_question():
    xml = to_xform((FIX / "amb_unsupported_expr.xlsx").read_bytes())
    with pytest.raises(UnsupportedExpression) as e:
        compile_spec(xml)
    assert "initial" in str(e.value) and "substr" in str(e.value)


def test_an_unsupported_choice_filter_names_its_question(xml):
    bad = xml.replace("[district= /data/district ]", "[substr(district, 1, 1)= 'p']")
    assert bad != xml
    with pytest.raises(UnsupportedExpression) as e:
        compile_spec(bad)
    assert "block" in str(e.value)


def test_computed_required_is_refused_not_dropped(xml):
    old = '<bind nodeset="/data/asha_name" type="string" required="true()"/>'
    new = '<bind nodeset="/data/asha_name" type="string" required=" /data/district  = \'x\'"/>'
    bad = xml.replace(old, new)
    assert bad != xml
    with pytest.raises(FormCompileError, match="asha_name"):
        compile_spec(bad)


def test_unknown_question_type_is_refused(xml):
    old = '<bind nodeset="/data/visit_location" type="geopoint"/>'
    bad = xml.replace(old, '<bind nodeset="/data/visit_location" type="geotrace"/>')
    assert bad != xml
    with pytest.raises(FormCompileError, match="visit_location"):
        compile_spec(bad)


def test_a_name_reused_in_a_nested_repeat_is_refused(xml):
    bad = xml.replace("<dose_ml/>", "<child_name/>")
    assert bad != xml
    with pytest.raises(FormCompileError) as e:
        compile_spec(bad)
    msg = str(e.value)
    assert "child_name" in msg
    assert "/data/children/child_name" in msg and "/data/children/ifa_doses/child_name" in msg


def test_a_name_reused_across_groups_is_refused(xml):
    bad = xml.replace("<asha_name/>", "<district/>", 1)
    assert bad != xml
    with pytest.raises(FormCompileError, match="district"):
        compile_spec(bad)


# ---- review fixes ---------------------------------------------------------

import openpyxl  # noqa: E402

from app.forms.compile import _short  # noqa: E402
from app.forms.expr import Prop, Ref, parse, parse_choice_filter  # noqa: E402


def test_the_emitted_choice_filter_parses_in_choice_filter_mode_only(spec):
    cf = find(spec["nodes"], "block")["choice_filter"]
    assert cf == "district = ${district}"
    tree = parse_choice_filter(cf)
    assert tree.left == Prop("district") and tree.right == Ref("district")
    with pytest.raises(UnsupportedExpression, match="bare name"):
        parse(cf)


def test_leftover_raw_paths_are_rejected_by_parse():
    for raw in ("../a/b", "/data", "/data/a", "../a", "a/b", "current()/x"):
        with pytest.raises(UnsupportedExpression):
            parse(raw)
        with pytest.raises(UnsupportedExpression):
            parse_choice_filter(raw)


def _single_language_xml():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "survey"
    ws.append(["type", "name", "label"])
    ws.append(["text", "who", "Name"])
    ws.append(["select_one yn", "ok", "OK?"])
    ws = wb.create_sheet("choices")
    ws.append(["list_name", "name", "label"])
    ws.append(["yn", "y", "Yes"])
    ws = wb.create_sheet("settings")
    ws.append(["form_title", "form_id", "version"])
    ws.append(["Solo", "solo", "1"])
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return to_xform(buf.getvalue())


def test_a_single_language_form_has_a_uniform_shape():
    spec = compile_spec(_single_language_xml())
    assert spec["languages"] == ["default"]
    assert spec["default_language"] == "default"
    assert spec["title"] == {"default": "Solo"}
    who = find(spec["nodes"], "who")
    assert set(who["label"]) <= set(spec["languages"])
    assert spec["choices"]["yn"][0]["label"] == {"default": "Yes"}


@pytest.mark.parametrize("raw, short", [
    ("hi (Hindi)", "hi"), ("Hindi (hi)", "hi"), ("en (English)", "en"),
    ("English (en)", "en"), ("hi", "hi"), ("pt-BR (Portuguese)", "pt-BR"),
    ("Hindi", "Hindi"),
])
def test_language_codes_are_shortened_in_either_convention(raw, short):
    assert _short(raw) == short


def test_two_translations_with_the_same_short_code_are_refused(xml):
    bad = xml.replace('lang="en (English)"', 'lang="Hindi (hi)"')
    assert bad != xml
    with pytest.raises(FormCompileError, match="hi"):
        compile_spec(bad)


def test_a_group_or_repeat_name_colliding_with_a_leaf_is_refused(xml):
    later = xml.replace("<children_count/>", "<children/>")
    earlier = xml.replace("<asha_name/>", "<children/>", 1)
    for bad in (later, earlier):
        assert bad != xml
        with pytest.raises(FormCompileError, match="children"):
            compile_spec(bad)


_TEXT_BIND = '<bind nodeset="/data/asha_name" type="string" required="true()"/>'
_TEXT_CTL = ('<input ref="/data/asha_name"><label ref="jr:itext(\'/data/asha_name:label\')"/>'
             "</input>")
_UPLOAD = ('<upload ref="/data/asha_name" mediatype="{m}"><label '
           'ref="jr:itext(\'/data/asha_name:label\')"/></upload>')


@pytest.mark.parametrize("bind_type, control, expected", [
    ("time", None, "time"), ("dateTime", None, "dateTime"), ("decimal", None, "decimal"),
    ("barcode", None, "barcode"),
    ("binary", "audio/*", "audio"), ("binary", "video/*", "video"),
    ("binary", "application/*", "file"),
])
def test_remaining_phase_1_types_are_mapped(xml, bind_type, control, expected):
    bad = xml.replace(_TEXT_BIND, _TEXT_BIND.replace('type="string"', f'type="{bind_type}"'))
    assert _TEXT_CTL in bad
    if control:
        bad = bad.replace(_TEXT_CTL, _UPLOAD.format(m=control))
    assert bad != xml
    assert find(compile_spec(bad)["nodes"], "asha_name")["type"] == expected


def test_a_missing_choice_list_is_refused_not_left_dangling(xml):
    bad = xml.replace('<instance id="sex">', '<instance id="sex_gone">')
    assert bad != xml
    with pytest.raises(FormCompileError, match="child_sex"):
        compile_spec(bad)


def _with_calculates(*rows):
    """The AMB fixture with extra top-level calculate rows (name, formula) appended."""
    import io

    from openpyxl import load_workbook

    wb = load_workbook(FIX / "amb_asha_visit.xlsx")
    ws = wb["survey"]
    header = [c.value for c in ws[1]]
    for name, formula, *rest in rows:
        row = [None] * len(header)
        row[header.index("type")] = rest[0] if rest else "calculate"
        row[header.index("name")] = name
        row[header.index("calculation")] = formula
        if len(rest) > 1:
            row[header.index("relevant")] = rest[1]
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return to_xform(buf.getvalue())


def test_a_calculate_cycle_is_refused_and_names_the_fields():
    # unseeded, so it would settle at NaN and publish silently without the guard
    with pytest.raises(FormCompileError, match="cyc_a -> cyc_b -> cyc_a"):
        compile_spec(_with_calculates(("cyc_a", "${cyc_b} + 1"), ("cyc_b", "${cyc_a} + 1")))


def test_a_calculate_that_refers_to_itself_is_refused():
    with pytest.raises(FormCompileError, match="selfie -> selfie"):
        compile_spec(_with_calculates(("selfie", "${selfie} + 1")))


def test_a_longer_calculate_cycle_is_refused():
    with pytest.raises(FormCompileError) as e:
        compile_spec(_with_calculates(("c1", "${c2}"), ("c2", "${c3}"), ("c3", "${c1}")))
    assert all(n in str(e.value) for n in ("c1", "c2", "c3"))


def test_calculates_that_chain_without_a_cycle_compile(spec):
    out = compile_spec(_with_calculates(("late", "${early} + 1"), ("early", "${household_children}")))
    assert find(out["nodes"], "late")["calculate"] == "${early} + 1"


def test_a_cycle_through_an_ancestor_relevant_is_refused():
    # x sits in a group whose relevant reads y, and y reads x
    with pytest.raises(FormCompileError, match="x -> y -> x"):
        compile_spec(_with_calculates(
            ("g", None, "begin group", "${y} > 0"),
            ("x", "1"),
            ("g_end", None, "end group"),
            ("y", "${x}"),
        ))


def test_a_calculate_gated_by_a_later_calculate_through_a_group_is_not_a_cycle():
    out = compile_spec(_with_calculates(
        ("g", None, "begin group", "${y} > 0"),
        ("x", "1"),
        ("g_end", None, "end group"),
        ("y", "${household_children}"),
    ))
    assert find(out["nodes"], "x")["calculate"] == "1"
